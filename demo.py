


我&#160;65&#160;岁了，有三个子女，他们都不愿意支付赡养费，我该怎么办？


import streamlit as st
from openai import OpenAI
import re
import os
import base64
import io
from docx import Document
from fpdf import FPDF
from dotenv import load_dotenv

load_dotenv()

try:
    raw_api_key = st.secrets["INTERN_API_KEY"]
    raw_base_url = st.secrets["INTERN_API_BASE"]
except Exception:
    raw_api_key = os.getenv('INTERN_API_KEY', '')
    raw_base_url = os.getenv('INTERN_API_BASE', '')

api_key = re.sub(r'\s+', '', raw_api_key)
base_url = raw_base_url.strip()

# 流式输出下，超时机制会被弱化，因为只要在一直输出就不会断
client = OpenAI(api_key=api_key, base_url=base_url, timeout=60.0)

st.set_page_config(page_title="银发法盾 - AI 法律助手", page_icon="⚖️", layout="wide")

st.title("⚖️ 银发法盾")
st.markdown("**专为中老年人打造的法律助手**，支持图片案情识别、类案对比与文书导出。")

# 侧边栏设置
with st.sidebar:
    st.header("系统设置")
    model_option = st.selectbox("选择生成模型 (建议选多模态模型)", 
        ["deepseek-v4.1-flash-1017", "intern-s2", "Atria-Dawn"], index=0)
    
    st.markdown("---")
    st.markdown("**使用指南：**\n1. 输入文字或上传案情截图\n2. 等待AI出具法律意见\n3. 一键下载Word/PDF")

# 分栏：左侧输入，右侧输出
col1, col2 = st.columns([1, 1.2])

with col1:
    st.subheader("📝 案情录入")
    user_prompt = st.text_area("文字描述案情：", height=150, placeholder="例如：我 65 岁了，有三个子女...")
    
    uploaded_file = st.file_uploader("或者上传案情图片/截图 (可选)", type=["jpg", "jpeg", "png"])
    if uploaded_file is not None:
        st.image(uploaded_file, caption="已上传案情图", width=300)

    # 功能开关
    st.markdown("**🛠️ 高级功能配置：**")
    enable_compare = st.checkbox("开启【类案对比与裁判规则】", value=True)
    enable_ocr = st.checkbox("开启【图片OCR自动提取】", value=True)
    
    analyze_btn = st.button("🚀 生成法律意见", type="primary", use_container_width=True)

with col2:
    st.subheader("📋 法律意见书")
    
    if analyze_btn:
        if not user_prompt.strip() and uploaded_file is None:
            st.error("⚠️ 请输入案情或上传图片")
        else:
            try:
                # 构建系统提示词
                system_content = "你是一位拥有20年经验的资深老龄权益律师。请用通俗易懂的语言解答，并严格按【案情定性】【法律分析】【维权路径】【行动建议】四部分输出。"
                if enable_compare:
                    system_content += "\n此外，请额外增加第五部分【类案对比】：总结1-2个最高人民法院或地方法院的真实相似判例，说明法院的裁判规则。"

                messages = [{"role": "system", "content": system_content}]
                
                # 处理多模态图片输入
                if uploaded_file is not None and enable_ocr:
                    base64_image = base64.b64encode(uploaded_file.read()).decode('utf-8')
                    content_list = [
                        {"type": "text", "text": f"请根据图片中的案情内容进行分析。我的附加说明：{user_prompt}"},
                        {"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{base64_image}"}}
                    ]
                    messages.append({"role": "user", "content": content_list})
                else:
                    messages.append({"role": "user", "content": user_prompt})

                # ================= 核心修改：流式输出 =================
                response = client.chat.completions.create(
                    model=model_option, 
                    messages=messages, 
                    temperature=0.5,
                    stream=True  # 开启流式输出
                )

                result_placeholder = st.empty()
                full_result = ""
                st.info("🔄 AI律师正在撰写意见，请稍候...")

                # 像打字机一样实时显示
                for chunk in response:
                    if chunk.choices[0].delta.content is not None:
                        full_result += chunk.choices[0].delta.content
                        result_placeholder.markdown(full_result + "▌")

                # 最终生成结束后，去掉闪烁的光标
                result_placeholder.markdown(full_result)
                st.success("✅ 分析完成！")
                
                # 把最终结果存下来，方便导出
                result = full_result
                # ====================================================
                
                # --- 文档导出功能 ---
                st.markdown("### 📥 导出法律意见书")
                export_col1, export_col2 = st.columns(2)
                
                # 导出 Word
                doc = Document()
                doc.add_heading('银发法盾 - AI法律意见书', 0)
                doc.add_paragraph(result)
                bio = io.BytesIO()
                doc.save(bio)
                export_col1.download_button(
                    label="下载 Word 版",
                    data=bio.getvalue(),
                    file_name="法律意见书.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    use_container_width=True
                )

                # 导出 PDF
                pdf = FPDF()
                pdf.add_page()
                pdf.set_font("Arial", size=12)
                pdf.multi_cell(0, 10, result.encode('latin-1', 'replace').decode('latin-1'))
                pdf_output = pdf.output(dest="S").encode("latin-1")
                
                export_col2.download_button(
                    label="下载 PDF 版 (英文/数字)",
                    data=pdf_output,
                    file_name="法律意见书.pdf",
                    mime="application/pdf",
                    use_container_width=True
                )
                st.info("💡 PDF包含中文时建议使用Word版，或前往打印店另存为PDF。")

            except Exception as e:
                st.error(f"❌ 发生错误：{str(e)}")
